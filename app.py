import os
import re
import pdfplumber
import pandas as pd
import streamlit as st
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Page Config
st.set_page_config(
    page_title="FinPulse | Bank Statement & OD Analytics", 
    page_icon="💳", 
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Premium Styling
st.markdown("""
    <style>
    .stApp { background-color: #F8FAFC; font-family: 'Inter', -apple-system, sans-serif; }
    .header-container {
        background: linear-gradient(135deg, #0F172A 0%, #1E293B 100%);
        padding: 24px 32px; border-radius: 16px; color: white; margin-bottom: 25px;
        box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.15);
    }
    .header-title { font-size: 28px !important; font-weight: 700 !important; color: #FFFFFF !important; margin: 0 !important; }
    .header-subtitle { font-size: 14px; color: #94A3B8; margin-top: 4px; }
    .metric-card {
        background: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 14px; padding: 18px 20px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.03); transition: transform 0.2s;
    }
    .metric-card:hover { transform: translateY(-2px); }
    .metric-label { font-size: 12px; font-weight: 600; text-transform: uppercase; color: #64748B; margin-bottom: 6px; }
    .metric-value { font-size: 22px; font-weight: 700; color: #0F172A; }
    .val-credit { color: #16A34A; }
    .val-debit { color: #DC2626; }
    .val-cash { color: #2563EB; }
    .val-charge { color: #D97706; }
    .val-interest { color: #7C3AED; }
    .stDownloadButton > button {
        width: 100%; background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%) !important;
        color: white !important; border: none !important; border-radius: 8px !important;
        padding: 10px 18px !important; font-weight: 600 !important; font-size: 14px !important;
    }
    #MainMenu {visibility: hidden;} footer {visibility: hidden;}
    </style>
""", unsafe_allow_html=True)

# Helper Functions
def clean_number(val):
    if not val or str(val).strip() == "" or str(val).lower() == "none":
        return 0.0
    val_str = str(val).replace(",", "").strip()
    is_negative = "-" in val_str
    match = re.search(r'\d+\.?\d*', val_str)
    if match:
        try:
            num = float(match.group(0))
            return -num if is_negative else num
        except ValueError:
            return 0.0
    return 0.0

def parse_bank_statement_pdf_perfect(pdf_file, password=None):
    all_rows = []
    pdf_kwargs = {'password': password} if password else {}

    with pdfplumber.open(pdf_file, **pdf_kwargs) as pdf:
        for page in pdf.pages:
            words = page.extract_words()
            lines = {}
            for w in words:
                top_key = round(w['top'], 1)
                matched_line = None
                for k in lines.keys():
                    if abs(k - top_key) < 3.5:
                        matched_line = k
                        break
                if matched_line is not None:
                    lines[matched_line].append(w)
                else:
                    lines[top_key] = [w]
            
            for k in sorted(lines.keys()):
                line_words = sorted(lines[k], key=lambda x: x['x0'])
                line_text = " ".join([w['text'] for w in line_words])
                
                if "Date" in line_text and "Closing Balance" in line_text:
                    continue
                if "Withdrawal" in line_text or "Deposit" in line_text or "Narration" in line_text:
                    continue
                    
                first_word = line_words[0]['text'] if line_words else ""
                if not re.match(r'^\d{2}/\d{2}/\d{2}$', first_word):
                    continue
                
                date_str = first_word
                narration_words, ref_words = [], []
                value_dt_str, withdrawal_str, deposit_str, balance_str = "", "", "", ""
                
                for w in line_words:
                    x0, txt = w['x0'], w['text']
                    if x0 < 60:
                        continue
                    elif 60 <= x0 < 220:
                        narration_words.append(txt)
                    elif 220 <= x0 < 340:
                        ref_words.append(txt)
                    elif 340 <= x0 < 400:
                        if re.match(r'^\d{2}/\d{2}/\d{2}$', txt):
                            value_dt_str = txt
                    elif 400 <= x0 < 470:
                        withdrawal_str = txt
                    elif 470 <= x0 < 540:
                        deposit_str = txt
                    elif x0 >= 540:
                        balance_str = txt
                
                all_rows.append({
                    "Date": date_str,
                    "Narration": " ".join(narration_words),
                    "Chq/Ref.No.": " ".join(ref_words),
                    "Value Dt": value_dt_str if value_dt_str else date_str,
                    "Withdrawal Amt": clean_number(withdrawal_str),
                    "Deposit Amt": clean_number(deposit_str),
                    "Closing Balance": clean_number(balance_str)
                })

    df = pd.DataFrame(all_rows)
    if not df.empty:
        # Original sequence tracking for accurate day-end balance selection
        df['Original_Order'] = df.index
    return df

def is_interest_entry(narration):
    narr = str(narration).upper()
    exclude_keywords = ['INT TRF', 'INT.TRF', 'INTERNAL', 'INTERCITY', 'INTEGRATED', 'TRANSFER', 'NEFT', 'RTGS', 'UPI']
    for ex in exclude_keywords:
        if ex in narr and 'DEBITED' not in narr and 'INTEREST' not in narr:
            return False
            
    interest_patterns = [
        r'INTEREST\s+DEBITED',
        r'INTEREST\s+CHARGED',
        r'INT\.?\s*COLL',
        r'OD\s+INTEREST',
        r'INTEREST\s+DRAWDOWN',
        r'COMPOUND\s+INTEREST',
        r'PENAL\s+INTEREST'
    ]
    for pattern in interest_patterns:
        if re.search(pattern, narr):
            return True
            
    if 'INTEREST' in narr and ('DEBIT' in narr or 'COLL' in narr or 'CHARGE' in narr):
        return True
        
    return False

def is_cash_deposit(narration):
    narr = str(narration).upper().replace(" ", "").replace("-", "")
    if "CASHFREE" in narr:
        return False
    if "CASH" in narr or "CDM" in narr or "CSH" in narr or "BYCASH" in narr:
        return True
    return False

def calculate_od_interest_multirate(df, od_limit, rate_slabs, start_date=None, end_date=None):
    df['Parsed_Date'] = pd.to_datetime(df['Value Dt'], format='%d/%m/%y', errors='coerce')
    df['Parsed_Date'] = df['Parsed_Date'].fillna(pd.to_datetime(df['Date'], format='%d/%m/%y', errors='coerce'))
    df = df.dropna(subset=['Parsed_Date']).reset_index(drop=True)
    
    # Filter Date Range
    if start_date and end_date:
        df = df[(df['Parsed_Date'].dt.date >= start_date) & (df['Parsed_Date'].dt.date <= end_date)]
    
    if df.empty:
        return df, pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    # Sort strictly by Parsed_Date and Original_Order to maintain exact transaction order
    df = df.sort_values(by=['Parsed_Date', 'Original_Order']).reset_index(drop=True)

    # 1. Exact Day-End Balance: Take the LAST entry of each date according to Original_Order
    idx_last = df.groupby('Parsed_Date')['Original_Order'].idxmax()
    daily_last = df.loc[idx_last, ['Parsed_Date', 'Closing Balance']].reset_index(drop=True)

    # 2. Continuous Date Range Fill
    full_date_range = pd.date_range(start=daily_last['Parsed_Date'].min(), end=daily_last['Parsed_Date'].max(), freq='D')
    daily_df = pd.DataFrame({'Parsed_Date': full_date_range})
    daily_df = pd.merge(daily_df, daily_last, on='Parsed_Date', how='left')
    daily_df['Closing Balance'] = daily_df['Closing Balance'].ffill()

    # 3. Dynamic Interest Rate Assignment Per Day
    def get_applicable_rate(row_date):
        c_date = row_date.date()
        for slab in rate_slabs:
            s_from = slab['from_date']
            s_to = slab['to_date']
            if s_from <= c_date <= s_to:
                return slab['rate']
        return rate_slabs[0]['rate'] if rate_slabs else 9.5

    daily_df['Applied_Rate'] = daily_df['Parsed_Date'].apply(get_applicable_rate)

    # 4. Utilized Amount & Daily Interest Calculation
    daily_df['Utilized_OD_Amount'] = daily_df['Closing Balance'].apply(lambda bal: abs(bal) if bal < 0 else 0.0)
    daily_df['Daily_Interest'] = daily_df['Utilized_OD_Amount'] * ((daily_df['Applied_Rate'] / 100.0) / 365.0)

    # 5. Monthly System Calculated Interest
    daily_df['Year_Month'] = daily_df['Parsed_Date'].dt.to_period('M')
    monthly_calc = daily_df.groupby('Year_Month')['Daily_Interest'].sum().reset_index()
    monthly_calc.rename(columns={'Daily_Interest': 'Calculated Interest (System)'}, inplace=True)

    # 6. Extract Actual Bank Charged Interest (Strict Match)
    df['Is_Interest_Entry'] = df['Narration'].apply(is_interest_entry)
    bank_int_df = df[df['Is_Interest_Entry'] & (df['Withdrawal Amt'] > 0)].copy()

    def get_target_month(row):
        narration = str(row['Narration']).upper()
        match = re.search(r'(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[a-zA-Z]*[\s\-]*(20\d{2}|\d{2})', narration)
        if match:
            month_str = match.group(1)
            year_str = match.group(2)
            if len(year_str) == 2:
                year_str = "20" + year_str
            dt_parsed = pd.to_datetime(f"01-{month_str}-{year_str}", format='%d-%b-%Y', errors='coerce')
            if not pd.isna(dt_parsed):
                return dt_parsed.to_period('M')
        return row['Parsed_Date'].to_period('M')

    if not bank_int_df.empty:
        bank_int_df['Year_Month'] = bank_int_df.apply(get_target_month, axis=1)
        monthly_bank = bank_int_df.groupby('Year_Month')['Withdrawal Amt'].sum().reset_index()
        monthly_bank.rename(columns={'Withdrawal Amt': 'Bank Charged Interest (Actual)'}, inplace=True)
    else:
        monthly_bank = pd.DataFrame(columns=['Year_Month', 'Bank Charged Interest (Actual)'])

    # 7. Merge & Reconciliation
    monthly_report = pd.merge(monthly_calc, monthly_bank, on='Year_Month', how='left').fillna(0.0)
    monthly_report['Difference (Excess/Short)'] = monthly_report['Bank Charged Interest (Actual)'] - monthly_report['Calculated Interest (System)']
    
    monthly_report['Month'] = monthly_report['Year_Month'].astype(str)
    monthly_report = monthly_report[['Month', 'Calculated Interest (System)', 'Bank Charged Interest (Actual)', 'Difference (Excess/Short)']]

    return df, daily_df, monthly_report, bank_int_df

def format_excel_sheet(workbook, sheet_name):
    ws = workbook[sheet_name]
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

        header_name = str(col[0].value)
        if any(amt_key in header_name for amt_key in ["Amt", "Balance", "Amount", "Interest", "Total", "Difference", "Rate"]):
            for cell in list(col)[1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '₹#,##0.00' if "Rate" not in header_name else '0.00%'
                    cell.alignment = Alignment(horizontal="right")

def convert_df_to_csv(df):
    return df.to_csv(index=False).encode('utf-8')

# Streamlit App UI
st.markdown("""
    <div class="header-container">
        <div class="header-title">💼 Financial Statement & OD Analytics</div>
        <div class="header-subtitle">Automated Bank Statement Audit, Multi-Rate Interest Computation & Category Exports</div>
    </div>
""", unsafe_allow_html=True)

st.sidebar.markdown("### ⚙️ Configuration")
uploaded_file = st.sidebar.file_uploader("Upload Bank Statement (PDF)", type=["pdf"])
pdf_password = st.sidebar.text_input("PDF Password (If Encrypted)", type="password")

st.sidebar.markdown("---")
st.sidebar.markdown("### 💰 Credit Limit")
od_limit = st.sidebar.number_input("Sanctioned OD Limit (₹)", value=15000000.0, step=100000.0, format="%.2f")

st.sidebar.markdown("---")
st.sidebar.markdown("### 📈 Multi-Date Interest Rates Slabs")

if 'num_slabs' not in st.session_state:
    st.session_state.num_slabs = 1

col_btn1, col_btn2 = st.sidebar.columns(2)
if col_btn1.button("➕ Add Slab"):
    st.session_state.num_slabs += 1
if col_btn2.button("➖ Remove") and st.session_state.num_slabs > 1:
    st.session_state.num_slabs -= 1

rate_slabs = []
for i in range(st.session_state.num_slabs):
    st.sidebar.markdown(f"**Rate Slab {i+1}**")
    s_col1, s_col2 = st.sidebar.columns(2)
    
    default_start = datetime(2025, 4, 1).date() if i == 0 else datetime(2025, 10, 1).date()
    default_end = datetime(2026, 3, 31).date()
    
    f_date = s_col1.date_input(f"From (Slab {i+1})", value=default_start, key=f"from_{i}")
    t_date = s_col2.date_input(f"To (Slab {i+1})", value=default_end, key=f"to_{i}")
    rate_val = st.sidebar.number_input(f"Interest Rate % p.a. (Slab {i+1})", value=9.5, step=0.1, format="%.2f", key=f"rate_{i}")
    
    rate_slabs.append({
        'from_date': f_date,
        'to_date': t_date,
        'rate': rate_val
    })

if uploaded_file is not None:
    try:
        raw_df = parse_bank_statement_pdf_perfect(uploaded_file, password=pdf_password if pdf_password else None)
        
        if raw_df.empty:
            st.error("⚠️ Unable to extract statement data. Check PDF format.")
        else:
            raw_df['Parsed_Date'] = pd.to_datetime(raw_df['Value Dt'], format='%d/%m/%y', errors='coerce')
            raw_df['Parsed_Date'] = raw_df['Parsed_Date'].fillna(pd.to_datetime(raw_df['Date'], format='%d/%m/%y', errors='coerce'))
            raw_df = raw_df.dropna(subset=['Parsed_Date']).reset_index(drop=True)
            
            min_date = raw_df['Parsed_Date'].min().date()
            max_date = raw_df['Parsed_Date'].max().date()
            
            st.sidebar.markdown("---")
            st.sidebar.markdown("### 📅 Overall Statement Period")
            date_range = st.sidebar.date_input("Filter Dates", value=(min_date, max_date), min_value=min_date, max_value=max_date)
            
            start_date, end_date = None, None
            if isinstance(date_range, tuple) and len(date_range) == 2:
                start_date, end_date = date_range[0], date_range[1]
            
            filtered_df, daily_summary, monthly_report, bank_int_df = calculate_od_interest_multirate(
                raw_df, od_limit, rate_slabs, start_date, end_date
            )
            
            if filtered_df.empty:
                st.warning("No transactions found in selected date range.")
            else:
                clean_export_df = filtered_df.drop(columns=['Parsed_Date', 'Is_Interest_Entry', 'Original_Order'], errors='ignore')
                deposits_df = clean_export_df[clean_export_df['Deposit Amt'] > 0]
                withdrawals_df = clean_export_df[clean_export_df['Withdrawal Amt'] > 0]
                
                filtered_df['Is_Cash_Deposit'] = filtered_df['Narration'].apply(is_cash_deposit)
                cash_df = filtered_df[filtered_df['Is_Cash_Deposit'] & (filtered_df['Deposit Amt'] > 0)].drop(columns=['Parsed_Date', 'Is_Interest_Entry', 'Is_Cash_Deposit', 'Original_Order'], errors='ignore')
                
                charge_mask = filtered_df['Narration'].str.contains(r'CHARGE|CHG|FEE|TAX|GST|COMMISSION|PENALTY', case=False, na=False)
                charges_df = filtered_df[charge_mask & (filtered_df['Withdrawal Amt'] > 0)].drop(columns=['Parsed_Date', 'Is_Interest_Entry', 'Is_Cash_Deposit', 'Original_Order'], errors='ignore')

                total_withdrawal = filtered_df['Withdrawal Amt'].sum()
                total_deposit = filtered_df['Deposit Amt'].sum()
                total_cash_deposit = cash_df['Deposit Amt'].sum() if not cash_df.empty else 0.0
                total_charges = charges_df['Withdrawal Amt'].sum() if not charges_df.empty else 0.0
                total_interest = daily_summary['Daily_Interest'].sum() if not daily_summary.empty else 0.0

                c1, c2, c3, c4, c5 = st.columns(5)
                c1.markdown(f'<div class="metric-card"><div class="metric-label">Total Deposits</div><div class="metric-value val-credit">₹{total_deposit:,.2f}</div></div>', unsafe_allow_html=True)
                c1.download_button("📥 Export Deposits", convert_df_to_csv(deposits_df), "Deposits.csv", "text/csv")

                c2.markdown(f'<div class="metric-card"><div class="metric-label">Total Debits</div><div class="metric-value val-debit">₹{total_withdrawal:,.2f}</div></div>', unsafe_allow_html=True)
                c2.download_button("📥 Export Debits", convert_df_to_csv(withdrawals_df), "Debits.csv", "text/csv")

                c3.markdown(f'<div class="metric-card"><div class="metric-label">Cash Deposited</div><div class="metric-value val-cash">₹{total_cash_deposit:,.2f}</div></div>', unsafe_allow_html=True)
                c3.download_button("📥 Export Cash", convert_df_to_csv(cash_df), "Cash_Deposits.csv", "text/csv")

                c4.markdown(f'<div class="metric-card"><div class="metric-label">Bank Charges</div><div class="metric-value val-charge">₹{total_charges:,.2f}</div></div>', unsafe_allow_html=True)
                c4.download_button("📥 Export Charges", convert_df_to_csv(charges_df), "Bank_Charges.csv", "text/csv")

                c5.markdown(f'<div class="metric-card"><div class="metric-label">OD Interest</div><div class="metric-value val-interest">₹{total_interest:,.2f}</div></div>', unsafe_allow_html=True)
                if not daily_summary.empty:
                    daily_exp = daily_summary[['Parsed_Date', 'Closing Balance', 'Utilized_OD_Amount', 'Applied_Rate', 'Daily_Interest']].copy()
                    daily_exp['Parsed_Date'] = daily_exp['Parsed_Date'].dt.strftime('%d/%m/%Y')
                    c5.download_button("📥 Export Daily Ledger", convert_df_to_csv(daily_exp), "OD_Daily_Ledger.csv", "text/csv")

                st.markdown("<br>", unsafe_allow_html=True)

                output_excel = "OD_Analysis_Report.xlsx"
                with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
                    clean_export_df.to_excel(writer, sheet_name='All_Transactions', index=False)
                    monthly_report.to_excel(writer, sheet_name='Monthly_Interest_Audit', index=False)
                    
                    if not bank_int_df.empty:
                        bank_int_export = bank_int_df.drop(columns=['Parsed_Date', 'Is_Interest_Entry', 'Is_Cash_Deposit', 'Year_Month', 'Original_Order'], errors='ignore')
                        bank_int_export.to_excel(writer, sheet_name='Bank_Interest_Entries', index=False)

                    if not daily_summary.empty:
                        summary_export = daily_summary[['Parsed_Date', 'Closing Balance', 'Utilized_OD_Amount', 'Applied_Rate', 'Daily_Interest']].copy()
                        summary_export['Parsed_Date'] = summary_export['Parsed_Date'].dt.strftime('%d/%m/%Y')
                        summary_export.columns = ['Date', 'Closing Balance', 'Utilized OD Amount', 'Applied Rate (%)', 'Daily Interest (INR)']
                        summary_export.to_excel(writer, sheet_name='Daily_OD_Ledger', index=False)

                    deposits_df.to_excel(writer, sheet_name='Deposits_Only', index=False)
                    withdrawals_df.to_excel(writer, sheet_name='Debits_Only', index=False)
                    cash_df.to_excel(writer, sheet_name='Cash_Deposits', index=False)
                    charges_df.to_excel(writer, sheet_name='Bank_Charges', index=False)
                    
                    wb = writer.book
                    for sheet in wb.sheetnames:
                        format_excel_sheet(wb, sheet)

                st.sidebar.markdown("---")
                with open(output_excel, "rb") as fp:
                    st.sidebar.download_button(
                        label="📊 Download Full Multi-Sheet Excel",
                        data=fp,
                        file_name="OD_Executive_Report.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

                tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
                    "📊 Monthly Interest Audit", 
                    "🏛️ Detected Bank Interest Entries",
                    "📋 Master Transactions", 
                    "📈 Daily OD Ledger", 
                    "💵 Cash Entries", 
                    "🧾 Bank Charges"
                ])
                
                with tab1:
                    st.markdown("### 🗓️ Monthly Interest Comparison & Reconciliation")
                    st.dataframe(
                        monthly_report.style.format({
                            "Calculated Interest (System)": "₹{:,.2f}",
                            "Bank Charged Interest (Actual)": "₹{:,.2f}",
                            "Difference (Excess/Short)": "₹{:,.2f}"
                        }), 
                        use_container_width=True
                    )
                    st.download_button(
                        "📥 Download Monthly Audit CSV", 
                        convert_df_to_csv(monthly_report), 
                        "Monthly_Interest_Audit.csv", 
                        "text/csv"
                    )

                with tab2:
                    st.markdown("### 🔍 Verified Bank Interest Debited Entries")
                    if not bank_int_df.empty:
                        disp_bank_int = bank_int_df.drop(columns=['Parsed_Date', 'Is_Interest_Entry', 'Is_Cash_Deposit', 'Year_Month', 'Original_Order'], errors='ignore')
                        st.dataframe(disp_bank_int, use_container_width=True)
                        st.download_button("📥 Download Bank Interest Entries CSV", convert_df_to_csv(disp_bank_int), "Bank_Interest_Entries.csv", "text/csv")
                    else:
                        st.info("No bank interest debit entries found.")

                with tab3:
                    st.dataframe(clean_export_df, use_container_width=True, height=420)
                    
                with tab4:
                    if not daily_summary.empty:
                        disp_daily = daily_summary[['Parsed_Date', 'Closing Balance', 'Utilized_OD_Amount', 'Applied_Rate', 'Daily_Interest']].copy()
                        disp_daily['Parsed_Date'] = disp_daily['Parsed_Date'].dt.strftime('%d/%m/%Y')
                        disp_daily.columns = ['Date', 'Closing Balance', 'Utilized OD Amount', 'Applied Rate (%)', 'Daily Interest (INR)']
                        st.dataframe(disp_daily, use_container_width=True, height=420)
                        
                with tab5:
                    st.markdown("### 💵 Cash Deposit Entries")
                    st.dataframe(cash_df, use_container_width=True, height=350)
                    st.download_button("📥 Download Cash Deposits CSV", convert_df_to_csv(cash_df), "Cash_Deposits.csv", "text/csv")
                    
                with tab6:
                    st.markdown("### 🧾 Detected Bank Charges & Fees")
                    st.dataframe(charges_df, use_container_width=True, height=350)
                    st.download_button("📥 Download Bank Charges CSV", convert_df_to_csv(charges_df), "Bank_Charges_Entries.csv", "text/csv")

    except Exception as e:
        st.error(f"Error processing document: {str(e)}")

else:
    st.info("👋 Upload a Bank Statement PDF to get started.")
